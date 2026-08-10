"""Tests for aggregated demand workflow: resolve_active_branch_excluded_sectors,
sector-branch mode, Road deduplication, and workbook/dummy consistency checks.

These tests exercise pure-function logic only — no real data files are read.
Integration tests covering actual ESTO/ninth projection data loading are out of
scope here because they require the full data directory.
"""
from __future__ import annotations

import inspect
import warnings

import pandas as pd
import pytest

import codebase.aggregated_demand_workflow as aggregated_demand_workflow
import codebase.supply_reconciliation.tables as supply_reconciliation_tables
from codebase.utilities.workflow_utils import clear_csv_cache, load_esto_csv, load_ninth_outlook_csv
from codebase.aggregated_demand_workflow import (
    _apply_first_projection_year_bridge,
    _extract_contextual_projection_years,
    _demand_branch_from_esto_flow,
    _find_nonenergy_base_year_fuel_gaps,
    _find_missing_nonenergy_template_fuels,
    _esto_flow_is_excluded,
    _sector_exclusion_suffix,
    build_aggregated_demand_as_dummy,
    resolve_active_branch_excluded_sectors,
)


def test_narrow_reference_loaders_reuse_shared_cache_without_mutating_source(tmp_path) -> None:
    clear_csv_cache()
    esto_source = tmp_path / "esto.csv"
    ninth_source = tmp_path / "ninth.csv"
    pd.DataFrame([
        {
            "economy": "01AUS", "flows": "14.01 Test", "products": "01 Fuel",
            "is_subtotal": False, "2022": 2.0, "unused": "ignore",
        },
    ]).to_csv(esto_source, index=False)
    pd.DataFrame([
        {
            "economy": "01_AUS", "scenarios": "reference", "sectors": "14_industry_sector",
            "sub1sectors": "x", "sub2sectors": "x", "sub3sectors": "x", "sub4sectors": "x",
            "fuels": "01_fuel", "subfuels": "x", "subtotal_results": False,
            "2022": 2.0, "2023": 3.0, "unused": "ignore",
        },
    ]).to_csv(ninth_source, index=False)

    esto_loaded = aggregated_demand_workflow._load_esto_base_csv(
        esto_source,
        economy="01_AUS",
        base_year=2022,
    )
    ninth_loaded = aggregated_demand_workflow._load_demand_csv(
        ninth_source,
        economy="01_AUS",
        final_year=2023,
    )
    cached_esto = load_esto_csv(
        esto_source,
        usecols=["economy", "flows", "products", "is_subtotal", "2022"],
    )
    cached_ninth = load_ninth_outlook_csv(
        ninth_source,
        usecols=[
            "economy", "scenarios", "sectors", "sub1sectors", "sub2sectors",
            "sub3sectors", "sub4sectors", "fuels", "subfuels", "subtotal_results", "2022", "2023",
        ],
    )

    assert esto_loaded["economy"].tolist() == ["01_AUS"]
    assert cached_esto["economy"].tolist() == ["01AUS"]
    assert "unused" not in esto_loaded.columns
    assert "unused" not in ninth_loaded.columns
    assert cached_ninth["economy"].tolist() == ["01_AUS"]
    clear_csv_cache()


def test_contextual_projection_uses_same_sector_esto_fuel_shares(monkeypatch):
    ninth = pd.DataFrame([
        {
            "economy": "20_USA", "scenarios": "reference",
            "sectors": "14_industry_sector", "sub1sectors": "14_02_construction",
            "sub2sectors": "x", "sub3sectors": "x", "sub4sectors": "x",
            "fuels": "01_x_thermal_coal", "subfuels": "x",
            "subtotal_results": False, "2023": 200.0,
        }
    ])
    esto = pd.DataFrame([
        {"economy": "20_USA", "flows": "14.02 Construction", "products": "p1", "is_subtotal": False, "2022": 25.0},
        {"economy": "20_USA", "flows": "14.02 Construction", "products": "p2", "is_subtotal": False, "2022": 75.0},
    ])
    mapping = pd.DataFrame([
        {"ninth_sector": "14_02_construction", "ninth_fuel": "01_x_thermal_coal", "esto_flow": "14.02 Construction", "esto_product": "p1"},
        {"ninth_sector": "14_02_construction", "ninth_fuel": "01_x_thermal_coal", "esto_flow": "14.02 Construction", "esto_product": "p2"},
    ])
    monkeypatch.setattr(
        aggregated_demand_workflow,
        "load_active_mapping_sheet",
        lambda sheet_name, workbook_path: mapping.copy(),
    )

    result, diagnostics = _extract_contextual_projection_years(
        ninth_df=ninth,
        esto_df=esto,
        csv_scenario="reference",
        esto_fuel_map={"p1": "Fuel A", "p2": "Fuel B"},
        base_year=2022,
        final_year=2023,
        exclude_own_use_td_losses=True,
    )

    values = result.set_index("leap_fuel_name")["value"].to_dict()
    assert values == pytest.approx({"Fuel A": 50.0, "Fuel B": 150.0})
    assert diagnostics.empty

    _, _, provenance = _extract_contextual_projection_years(
        ninth_df=ninth,
        esto_df=esto,
        csv_scenario="reference",
        esto_fuel_map={"p1": "Fuel A", "p2": "Fuel B"},
        base_year=2022,
        final_year=2023,
        exclude_own_use_td_losses=True,
        return_allocation_provenance=True,
    )
    assert set(provenance["allocation_method"]) == {"proportional_esto_base_year"}
    shares = provenance.set_index("esto_product")["share"].to_dict()
    assert shares == pytest.approx({"p1": 0.25, "p2": 0.75})


def test_first_projection_year_bridge_applies_same_offset_to_all_projection_years():
    raw = pd.DataFrame(
        [
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2022,
                "value": 100.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2023,
                "value": 160.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2024,
                "value": 170.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2025,
                "value": 0.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Gas",
                "year": 2022,
                "value": 50.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Gas",
                "year": 2023,
                "value": 40.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Gas",
                "year": 2024,
                "value": 45.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Hydrogen",
                "year": 2022,
                "value": 0.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Hydrogen",
                "year": 2023,
                "value": 30.0,
            },
        ]
    )

    bridged = _apply_first_projection_year_bridge(
        raw,
        base_year=2022,
        projection_start_year=2023,
        blend_weight=0.0,
        enabled=True,
    )

    values = (
        bridged.set_index(["leap_fuel_name", "year"])["value"]
        .to_dict()
    )
    assert values[("Electricity", 2023)] == pytest.approx(100.0)
    assert values[("Electricity", 2024)] == pytest.approx(110.0)
    assert values[("Electricity", 2025)] == pytest.approx(0.0)
    assert values[("Gas", 2023)] == pytest.approx(50.0)
    assert values[("Gas", 2024)] == pytest.approx(55.0)
    assert values[("Hydrogen", 2023)] == pytest.approx(30.0)


def test_first_projection_year_bridge_also_closes_first_year_drops():
    raw = pd.DataFrame(
        [
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2022,
                "value": 100.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2023,
                "value": 80.0,
            },
            {
                "economy": "20_USA",
                "scenario": "Reference",
                "leap_fuel_name": "Electricity",
                "year": 2024,
                "value": 90.0,
            },
        ]
    )

    bridged = _apply_first_projection_year_bridge(
        raw,
        base_year=2022,
        projection_start_year=2023,
        blend_weight=0.0,
        enabled=True,
    )

    values = bridged.set_index("year")["value"].to_dict()
    assert values[2022] == pytest.approx(100.0)
    assert values[2023] == pytest.approx(100.0)
    assert values[2024] == pytest.approx(110.0)


def test_build_aggregated_demand_bridge_mode_switch(monkeypatch):
    esto = pd.DataFrame(
        [
            {
                "economy": "20_USA",
                "flows": "14.02 Construction",
                "products": "p1",
                "is_subtotal": False,
                "2022": 100.0,
            }
        ]
    )
    projected = pd.DataFrame(
        [
            {
                "economy": "20_USA",
                "leap_fuel_name": "Fuel A",
                "year": 2023,
                "value": 160.0,
            },
            {
                "economy": "20_USA",
                "leap_fuel_name": "Fuel A",
                "year": 2024,
                "value": 170.0,
            }
        ]
    )
    monkeypatch.setattr(
        aggregated_demand_workflow,
        "_load_esto_base_csv",
        lambda *args, **kwargs: esto.copy(),
    )
    monkeypatch.setattr(
        aggregated_demand_workflow,
        "_load_demand_csv",
        lambda *args, **kwargs: pd.DataFrame(),
    )
    monkeypatch.setattr(
        aggregated_demand_workflow,
        "load_fuel_mapping",
        lambda *args, **kwargs: {"p1": "Fuel A"},
    )
    monkeypatch.setattr(
        aggregated_demand_workflow,
        "_extract_contextual_projection_years",
        lambda *args, **kwargs: (projected.copy(), pd.DataFrame()),
    )

    off = aggregated_demand_workflow.build_aggregated_demand(
        economy="20_USA",
        scenario="Reference",
        base_year=2022,
        final_year=2024,
        data_path="unused.csv",
        esto_data_path="unused.csv",
        fuel_mappings_path="unused.xlsx",
        apply_first_projection_year_bridge=False,
    )
    on = aggregated_demand_workflow.build_aggregated_demand(
        economy="20_USA",
        scenario="Reference",
        base_year=2022,
        final_year=2024,
        data_path="unused.csv",
        esto_data_path="unused.csv",
        fuel_mappings_path="unused.xlsx",
        apply_first_projection_year_bridge=True,
    )

    off_values = off.set_index("year")["value"].to_dict()
    on_values = on.set_index("year")["value"].to_dict()
    assert off_values[2022] == pytest.approx(100.0)
    assert off_values[2023] == pytest.approx(160.0)
    assert off_values[2024] == pytest.approx(170.0)
    assert on_values[2022] == pytest.approx(100.0)
    assert on_values[2023] == pytest.approx(100.0)
    assert on_values[2024] == pytest.approx(110.0)
from codebase.supply_reconciliation.config import (
    DETAILED_DEMAND_BRANCHES_ACTIVE,
    LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP,
)


# ── resolve_active_branch_excluded_sectors ────────────────────────────────────


class TestResolveActiveBranchExcludedSectors:
    """Unit tests for the resolve_active_branch_excluded_sectors helper."""

    SECTOR_MAP = LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP

    def test_none_inputs_return_none(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=None,
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is None

    def test_empty_active_branches_returns_base_only(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=[],
            sector_map=self.SECTOR_MAP,
            base_excluded=["14_industry_sector"],
        )
        assert result == ["14_industry_sector"]

    def test_active_branches_without_base_returns_sectors(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Industry"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result == ["14_industry_sector"]

    def test_base_excluded_preserved_in_output_order(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Industry"],
            sector_map=self.SECTOR_MAP,
            base_excluded=["16_01_buildings"],
        )
        # base comes first, then active-branch additions
        assert result == ["16_01_buildings", "14_industry_sector"]

    def test_road_excludes_road_sector(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Road"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is not None
        assert result.count("15_02_road") == 1

    def test_road_deduplication_when_base_also_contains_road(self):
        """If 15_02_road is already in base_excluded, it must not be duplicated
        when Road is also active."""
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Road"],
            sector_map=self.SECTOR_MAP,
            base_excluded=["15_02_road"],
        )
        assert result is not None
        assert result.count("15_02_road") == 1

    def test_multiple_branches_deduplication(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Road", "Industry", "Buildings"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is not None
        assert len(result) == len(set(result)), "Result must not contain duplicate sector codes"
        assert "15_02_road" in result
        assert "14_industry_sector" in result
        assert "16_01_buildings" in result
        assert result.count("15_02_road") == 1

    def test_transport_non_road_excludes_expected_sectors(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Transport non road"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is not None
        expected = {
            "15_01_domestic_air_transport",
            "15_03_rail",
            "15_04_domestic_navigation",
            "15_05_pipeline_transport",
            "15_06_nonspecified_transport",
        }
        assert set(result) == expected

    def test_international_transport_excludes_bunkers(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["International transport"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert set(result) == {
            "04_international_marine_bunkers",
            "05_international_aviation_bunkers",
        }

    def test_other_sector_excludes_agriculture_and_fishing_and_nonspecified(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Other sector"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is not None
        assert "16_02_agriculture_and_fishing" in result
        assert "16_05_nonspecified_others" in result
        assert "16_01_buildings" not in result

    def test_buildings_excludes_only_buildings_sub1sector(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Buildings"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result == ["16_01_buildings"]

    def test_unknown_branch_warns_and_continues(self):
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            result = resolve_active_branch_excluded_sectors(
                active_branches=["Unknown branch name"],
                sector_map=self.SECTOR_MAP,
                base_excluded=["14_industry_sector"],
            )
        # Should still return the base_excluded even though the branch was unknown
        assert result == ["14_industry_sector"]
        assert len(caught) == 1
        assert "Unknown branch name" in str(caught[0].message)

    def test_returns_none_when_all_empty(self):
        result = resolve_active_branch_excluded_sectors(
            active_branches=[],
            sector_map=self.SECTOR_MAP,
            base_excluded=[],
        )
        assert result is None

    def test_order_is_base_first_then_active(self):
        """base_excluded entries should appear before active-branch entries."""
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Industry"],
            sector_map=self.SECTOR_MAP,
            base_excluded=["16_01_buildings", "16_02_agriculture_and_fishing"],
        )
        assert result is not None
        idx_buildings = result.index("16_01_buildings")
        idx_industry = result.index("14_industry_sector")
        assert idx_buildings < idx_industry


# ── Config structure validation ───────────────────────────────────────────────


class TestLeapDemandGroupEstoSectorMapConfig:
    """Validate that the config mapping has the expected structure and known keys."""

    def test_map_contains_all_expected_groups(self):
        expected_groups = {
            "Road",
            "Transport non road",
            "International transport",
            "Industry",
            "Other sector",
            "Non Energy Use",
            "Buildings",
        }
        assert set(LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP.keys()) == expected_groups

    def test_road_maps_to_15_02_road(self):
        assert LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["Road"] == ["15_02_road"]

    def test_industry_maps_to_top_level_sector(self):
        assert "14_industry_sector" in LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["Industry"]

    def test_buildings_maps_to_16_01(self):
        assert "16_01_buildings" in LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["Buildings"]
        # Must not accidentally include other 16-sector codes
        assert all(
            code.startswith("16_01") for code in LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["Buildings"]
        )

    def test_other_sector_includes_agriculture_fishing_and_nonspecified(self):
        other_codes = set(LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["Other sector"])
        assert "16_02_agriculture_and_fishing" in other_codes
        assert "16_05_nonspecified_others" in other_codes
        assert "17_nonenergy_use" not in other_codes
        # Should not accidentally include buildings
        assert "16_01_buildings" not in other_codes

    def test_nonenergy_use_has_its_own_group(self):
        assert LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["Non Energy Use"] == ["17_nonenergy_use"]

    def test_international_transport_includes_bunkers(self):
        international = set(LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP["International transport"])
        assert "04_international_marine_bunkers" in international
        assert "05_international_aviation_bunkers" in international
        assert "15_02_road" not in international

    def test_all_esto_codes_are_valid_strings(self):
        for group, codes in LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP.items():
            for code in codes:
                assert isinstance(code, str) and code, (
                    f"Empty or non-string code in group '{group}'"
                )

    def test_default_active_branches_is_none(self):
        assert DETAILED_DEMAND_BRANCHES_ACTIVE is None


# ── _esto_flow_is_excluded ────────────────────────────────────────────────────


class TestEstoFlowIsExcluded:
    """Verify that the prefix-matching logic handles all sector codes in LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP."""

    def _excluded(self, code: str, flow: str) -> bool:
        return _esto_flow_is_excluded(flow, {code})

    def test_14_industry_sector_excludes_all_14_flows(self):
        assert self._excluded("14_industry_sector", "14.01 Mining")
        assert self._excluded("14_industry_sector", "14.03 Manufacturing")
        assert not self._excluded("14_industry_sector", "15.02 Road")

    def test_15_02_road_excludes_road_sub_sectors(self):
        assert self._excluded("15_02_road", "15.02 Road")
        assert self._excluded("15_02_road", "15.02.01 Passenger")
        assert self._excluded("15_02_road", "15.02.02 Freight")
        assert not self._excluded("15_02_road", "15.01 Air")
        assert not self._excluded("15_02_road", "15.03 Rail")

    def test_16_01_buildings_excludes_commercial_and_residential(self):
        assert self._excluded("16_01_buildings", "16.01 Buildings")
        assert self._excluded("16_01_buildings", "16.01.01 Commercial and public services")
        assert self._excluded("16_01_buildings", "16.01.02 Residential")
        assert not self._excluded("16_01_buildings", "16.02 Agriculture")

    def test_16_02_agriculture_and_fishing_excludes_agriculture_sub_flows(self):
        assert self._excluded("16_02_agriculture_and_fishing", "16.02 Agriculture and fishing")
        assert self._excluded("16_02_agriculture_and_fishing", "16.02.03 Agriculture")
        assert self._excluded("16_02_agriculture_and_fishing", "16.02.04 Fishing")
        assert not self._excluded("16_02_agriculture_and_fishing", "16.05 Non-specified")

    def test_04_and_05_bunkers_excluded(self):
        assert self._excluded("04_international_marine_bunkers", "04 International marine bunkers")
        assert self._excluded("05_international_aviation_bunkers", "05 International aviation bunkers")

    def test_transport_non_road_sectors_excluded(self):
        for code in [
            "15_01_domestic_air_transport",
            "15_03_rail",
            "15_04_domestic_navigation",
            "15_05_pipeline_transport",
            "15_06_nonspecified_transport",
        ]:
            prefix_num = code.split("_")[0] + "." + code.split("_")[1]
            assert self._excluded(code, prefix_num + " some label"), (
                f"{code} should exclude flows starting with {prefix_num}"
            )

    def test_road_not_excluded_by_non_road_codes(self):
        for code in [
            "15_01_domestic_air_transport",
            "15_03_rail",
            "15_04_domestic_navigation",
        ]:
            assert not self._excluded(code, "15.02 Road"), (
                f"{code} must not exclude road flows"
            )


# ── _sector_exclusion_suffix ──────────────────────────────────────────────────


class TestSectorExclusionSuffix:
    """Verify filenames encode the correct short codes for active-branch exclusions."""

    def test_no_exclusions_gives_empty_suffix(self):
        assert _sector_exclusion_suffix(None) == ""
        assert _sector_exclusion_suffix([]) == ""

    def test_industry_sector_gives_ind_suffix(self):
        suffix = _sector_exclusion_suffix(["14_industry_sector"])
        assert suffix == "_no_IND"

    def test_road_gives_rd_suffix(self):
        suffix = _sector_exclusion_suffix(["15_02_road"])
        assert "RD" in suffix

    def test_buildings_gives_bd_suffix(self):
        suffix = _sector_exclusion_suffix(["16_01_buildings"])
        assert "BD" in suffix

    def test_effective_exclusions_from_active_branches_produce_deterministic_suffix(self):
        """Regardless of input order, suffix should be deterministic."""
        excl_a = resolve_active_branch_excluded_sectors(
            active_branches=["Industry", "Buildings"],
            sector_map=LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP,
        )
        excl_b = resolve_active_branch_excluded_sectors(
            active_branches=["Buildings", "Industry"],
            sector_map=LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP,
        )
        assert _sector_exclusion_suffix(excl_a) == _sector_exclusion_suffix(excl_b)

    def test_road_produces_rd_in_suffix(self):
        excl = resolve_active_branch_excluded_sectors(
            active_branches=["Road"],
            sector_map=LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP,
        )
        suffix = _sector_exclusion_suffix(excl)
        assert suffix.count("RD") == 1


# ── build_aggregated_demand_as_dummy signature ────────────────────────────────


class TestBuildAggregatedDemandAsDummySignature:
    """Verify the function signature accepts the new use_sector_branches parameter."""

    def test_accepts_use_sector_branches_parameter(self):
        sig = inspect.signature(build_aggregated_demand_as_dummy)
        assert "use_sector_branches" in sig.parameters

    def test_use_sector_branches_defaults_to_false(self):
        sig = inspect.signature(build_aggregated_demand_as_dummy)
        param = sig.parameters["use_sector_branches"]
        assert param.default is False

    def test_accepts_excluded_sectors_parameter(self):
        sig = inspect.signature(build_aggregated_demand_as_dummy)
        assert "excluded_sectors" in sig.parameters

    def test_excluded_sectors_defaults_to_none(self):
        sig = inspect.signature(build_aggregated_demand_as_dummy)
        param = sig.parameters["excluded_sectors"]
        assert param.default is None


class TestBuildAggregatedDemandAsDummyForwarding:
    """Verify the dummy-demand helper forwards sector and own-use settings."""

    def test_forwards_sector_and_own_use_flags(self, monkeypatch):
        captured: dict[str, object] = {}

        def fake_build_aggregated_demand_all_scenarios(**kwargs):
            captured.update(kwargs)
            return pd.DataFrame(
                [
                    {
                        "economy": "20_USA",
                        "scenario": "Reference",
                        "leap_fuel_name": "Electricity",
                        "year": 2022,
                        "value": 10.0,
                        "sector": "14_industry_sector",
                    }
                ]
            )

        def fake_load_active_mapping_sheet(sheet_name, path):
            return pd.DataFrame(
                [
                    {
                        "raw_leap_fuel_name": "Electricity",
                        "esto_product": "17_electricity",
                    }
                ]
            )

        monkeypatch.setattr(
            aggregated_demand_workflow,
            "build_aggregated_demand_all_scenarios",
            fake_build_aggregated_demand_all_scenarios,
        )
        monkeypatch.setattr(
            aggregated_demand_workflow,
            "load_active_mapping_sheet",
            fake_load_active_mapping_sheet,
        )

        result = aggregated_demand_workflow.build_aggregated_demand_as_dummy(
            economy="20_USA",
            scenarios=["Reference"],
            excluded_sectors=["14_industry_sector"],
            exclude_own_use_td_losses=True,
            use_sector_branches=True,
        )

        assert captured["use_sector_branches"] is True
        assert captured["exclude_own_use_td_losses"] is True
        assert captured["excluded_sectors"] == ["14_industry_sector"]
        assert not result.empty
        assert list(result.columns) == [
            "economy",
            "scenario",
            "esto_product",
            "year",
            "demand_value",
            "demand_source",
        ]
        assert result.iloc[0]["demand_source"] == "aggregated_demand_projection"


class TestAggregatedDemandWorkbookModes:
    """Verify the exported branch structure changes with the sector toggle."""

    def test_flat_and_sector_split_branch_paths_differ(self, tmp_path):
        flat_path = tmp_path / "flat.xlsx"
        sector_path = tmp_path / "sector.xlsx"

        flat_df = pd.DataFrame(
            [
                {
                    "economy": "20_USA",
                    "scenario": "Reference",
                    "leap_fuel_name": "Electricity",
                    "year": 2022,
                    "value": 10.0,
                }
            ]
        )
        sector_df = pd.DataFrame(
            [
                {
                    "economy": "20_USA",
                    "scenario": "Reference",
                    "sector": "14_industry_sector",
                    "leap_fuel_name": "Electricity",
                    "year": 2022,
                    "value": 10.0,
                }
            ]
        )

        aggregated_demand_workflow.save_to_leap_export(
            flat_df,
            output_path=flat_path,
            use_sector_branches=False,
        )
        aggregated_demand_workflow.save_to_leap_export(
            sector_df,
            output_path=sector_path,
            use_sector_branches=True,
        )

        flat_export = pd.read_excel(flat_path, sheet_name="Export", header=None)
        sector_export = pd.read_excel(sector_path, sheet_name="Export", header=None)

        flat_branch_paths = set(flat_export.iloc[:, 0].astype(str))
        sector_branch_paths = set(sector_export.iloc[:, 0].astype(str))

        assert "Demand\\All demand aggregated\\Electricity" in flat_branch_paths
        assert "Demand\\All demand aggregated\\Industry\\Electricity" in sector_branch_paths
        assert "Demand\\All demand aggregated\\Industry\\Electricity" not in flat_branch_paths

    @pytest.mark.parametrize(
        ("flow", "expected_branch"),
        [
            ("15.02 Road", "Road"),
            ("04 International marine bunkers", "International transport"),
            ("05 International aviation bunkers", "International transport"),
            ("15.01 Domestic aviation", "Transport non road"),
            ("14.03 Manufacturing", "Industry"),
            ("16.01 Commercial and public services", "Buildings"),
            ("16.02 Residential", "Buildings"),
            ("16.03 Agriculture", "Other sector"),
            ("17 Non-energy use", "Non Energy Use"),
        ],
    )
    def test_requested_sector_branches_classify_source_flows(self, flow, expected_branch):
        assert _demand_branch_from_esto_flow(flow) == expected_branch

    def test_base_year_separates_nonenergy_from_other_sector(self):
        source = pd.DataFrame(
            [
                {
                    "economy": "01_AUS",
                    "flows": "16.05 Non-specified others",
                    "products": "07.17 Other products",
                    "is_subtotal": False,
                    "2022": 0.000641,
                },
                {
                    "economy": "01_AUS",
                    "flows": "17 Non-energy use",
                    "products": "07.17 Other products",
                    "is_subtotal": False,
                    "2022": 78.873358,
                },
            ]
        )

        result = aggregated_demand_workflow._extract_base_year(
            source,
            base_year=2022,
            exclude_own_use_td_losses=True,
            use_sector_branches=True,
        )
        values = result.groupby("sector")["value"].sum().to_dict()

        assert values == pytest.approx(
            {"Other sector": 0.000641, "Non Energy Use": 78.873358}
        )

    def test_nonenergy_fuel_coverage_finds_unmapped_and_missing_fuels(self):
        base_rows = pd.DataFrame(
            [
                {"sector": "Non Energy Use", "fuel_code": "mapped", "value": 2.0},
                {"sector": "Non Energy Use", "fuel_code": "unmapped", "value": 3.0},
                {"sector": "Other sector", "fuel_code": "other", "value": 4.0},
            ]
        )
        base_result = pd.DataFrame(
            [
                {
                    "sector": "Non Energy Use",
                    "leap_fuel_name": "Different fuel",
                    "value": 2.0,
                }
            ]
        )

        unmapped, missing = _find_nonenergy_base_year_fuel_gaps(
            base_rows=base_rows,
            base_result=base_result,
            esto_fuel_map={"mapped": "Expected fuel"},
        )

        assert unmapped == ["unmapped"]
        assert missing == ["Expected fuel"]

    def test_nonenergy_template_coverage_reports_only_nonzero_missing_fuels(self):
        demand = pd.DataFrame(
            [
                {
                    "scenario": "Reference",
                    "sector": "Non Energy Use",
                    "leap_fuel_name": "Naphtha",
                    "value": 2.0,
                },
                {
                    "scenario": "Reference",
                    "sector": "Non Energy Use",
                    "leap_fuel_name": "Biogas",
                    "value": 1.0,
                },
                {
                    "scenario": "Reference",
                    "sector": "Non Energy Use",
                    "leap_fuel_name": "Unused",
                    "value": 0.0,
                },
                {
                    "scenario": "Reference",
                    "sector": "Other sector",
                    "leap_fuel_name": "Coal",
                    "value": 3.0,
                },
            ]
        )
        template_paths = {
            r"Demand\All demand aggregated\Non Energy Use\Naphtha",
        }

        missing = _find_missing_nonenergy_template_fuels(
            demand=demand,
            scenarios=["Reference"],
            template_branch_paths=template_paths,
        )

        assert missing == ["Biogas"]


class TestReconciliationDemandInference:
    """Verify aggregated-demand reconciliation uses LEAP results to identify active branches."""

    def test_infers_active_branches_from_sector_table_and_deduplicates_road(self, monkeypatch):
        captured: dict[str, object] = {}

        monkeypatch.setattr(supply_reconciliation_tables, "USE_AGGREGATED_DEMAND_AS_DUMMY", True)
        monkeypatch.setattr(supply_reconciliation_tables, "DETAILED_DEMAND_BRANCHES_ACTIVE", None)
        monkeypatch.setattr(
            supply_reconciliation_tables,
            "AGGREGATED_DEMAND_EXCLUDED_SECTORS",
            None,
        )
        monkeypatch.setattr(
            supply_reconciliation_tables,
            "AGGREGATED_DEMAND_EXCLUDE_OWN_USE_TD_LOSSES",
            True,
        )
        monkeypatch.setattr(aggregated_demand_workflow, "USE_SECTOR_BRANCHES", True)

        def fake_load_results_sector_demand_table(**kwargs):
            return pd.DataFrame(
                [
                    {"sheet": "Industry", "demand_value": 10.0},
                    {"sheet": "Road", "demand_value": 15.0},
                ]
            )

        def fake_build_aggregated_demand_as_dummy(**kwargs):
            captured.update(kwargs)
            return pd.DataFrame(
                [
                    {
                        "economy": "20_USA",
                        "scenario": "Reference",
                        "esto_product": "17_electricity",
                        "year": 2022,
                        "demand_value": 10.0,
                        "demand_source": "aggregated_demand_projection",
                    }
                ]
            )

        monkeypatch.setattr(
            supply_reconciliation_tables,
            "load_results_sector_demand_table",
            fake_load_results_sector_demand_table,
        )
        monkeypatch.setattr(
            aggregated_demand_workflow,
            "build_aggregated_demand_as_dummy",
            fake_build_aggregated_demand_as_dummy,
        )

        result = supply_reconciliation_tables.load_results_demand_table(
            economies=["20_USA"],
            comparison_long_df=pd.DataFrame(),
            mapping_status_df=pd.DataFrame(),
        )

        assert not result.empty
        assert captured["use_sector_branches"] is True
        assert captured["exclude_own_use_td_losses"] is True
        assert captured["excluded_sectors"].count("15_02_road") == 1
        assert "14_industry_sector" in captured["excluded_sectors"]
        assert "15_02_road" in captured["excluded_sectors"]


# ── Reconciliation accounting logic ──────────────────────────────────────────


class TestReconciliationAccountingLogic:
    """
    Verify key invariants of the accounting design without running data pipelines.

    The intended accounting is:
        aggregated placeholder = total 9th/ESTO demand
                                 - 9th/ESTO demand for sectors with detailed LEAP branches

    Subtraction values come from 9th/ESTO source data, not LEAP result values.
    """

    SECTOR_MAP = LEAP_DEMAND_GROUP_ESTO_SECTOR_MAP

    def test_all_active_branches_produce_non_empty_exclusion_list(self):
        all_groups = list(self.SECTOR_MAP.keys())
        result = resolve_active_branch_excluded_sectors(
            active_branches=all_groups,
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is not None and len(result) > 0

    def test_no_active_branches_produces_full_aggregated_demand(self):
        """When no detailed branches are active, excluded sectors should be None
        (so the aggregated placeholder represents all demand)."""
        result = resolve_active_branch_excluded_sectors(
            active_branches=None,
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is None

    def test_road_active_produces_road_exclusion(self):
        excluded = resolve_active_branch_excluded_sectors(
            active_branches=["Road"],
            sector_map=self.SECTOR_MAP,
        )
        assert excluded == ["15_02_road"]

    def test_inactive_branches_not_excluded(self):
        """When a branch group is not in active_branches, its ESTO sectors must remain
        in the aggregated placeholder (i.e., not in the exclusion list)."""
        # Only Industry is active
        result = resolve_active_branch_excluded_sectors(
            active_branches=["Industry"],
            sector_map=self.SECTOR_MAP,
            base_excluded=None,
        )
        assert result is not None
        # Buildings and transport sectors should NOT be excluded
        assert "16_01_buildings" not in result
        assert "15_02_road" not in result
        assert "15_03_rail" not in result

    def test_exclusion_list_grows_monotonically_as_branches_are_added(self):
        """As more detailed branches become active, the exclusion list can only grow
        (or stay equal due to deduplication), never shrink."""
        excl_0 = resolve_active_branch_excluded_sectors(
            active_branches=[],
            sector_map=self.SECTOR_MAP,
        )
        excl_1 = resolve_active_branch_excluded_sectors(
            active_branches=["Industry"],
            sector_map=self.SECTOR_MAP,
        )
        excl_2 = resolve_active_branch_excluded_sectors(
            active_branches=["Industry", "Buildings"],
            sector_map=self.SECTOR_MAP,
        )
        excl_3 = resolve_active_branch_excluded_sectors(
            active_branches=["Industry", "Buildings", "Other sector"],
            sector_map=self.SECTOR_MAP,
        )
        len_0 = len(excl_0 or [])
        len_1 = len(excl_1 or [])
        len_2 = len(excl_2 or [])
        len_3 = len(excl_3 or [])
        assert len_0 <= len_1 <= len_2 <= len_3

    def test_road_branch_exclusion_is_15_02(self):
        excluded = resolve_active_branch_excluded_sectors(
            active_branches=["Road"],
            sector_map=self.SECTOR_MAP,
        )
        assert excluded == ["15_02_road"]


class TestAllZeroDemandBranchExportFilter:
    """All-zero branches clear existing template branches but do not create new ones."""

    @staticmethod
    def _row(
        sector: str,
        fuel: str,
        scenario: str,
        year: int,
        value: float,
    ) -> dict[str, object]:
        return {
            "economy": "20_USA",
            "sector": sector,
            "leap_fuel_name": fuel,
            "scenario": scenario,
            "year": year,
            "value": value,
        }

    def test_all_zero_branch_is_omitted_at_export_boundary(self, tmp_path):
        demand = pd.DataFrame(
            [
                self._row("Road", "Wind", "Current Accounts", 2022, 0.0),
                self._row("Road", "Wind", "Reference", 2023, 0.0),
                self._row("Road", "Wind", "Target", 2023, 5e-13),
                self._row("Road", "Electricity", "Current Accounts", 2022, 2.0),
                self._row("Road", "Electricity", "Reference", 2023, 3.0),
                self._row("Road", "Electricity", "Target", 2023, 4.0),
            ]
        )
        original = demand.copy(deep=True)
        output_path = tmp_path / "aggregated_demand.xlsx"

        result = aggregated_demand_workflow.save_aggregated_demand_as_leap_workbook(
            economy="20_USA",
            output_path=output_path,
            scenarios=["Current Accounts", "Reference", "Target"],
            id_lookup_path=None,
            use_sector_branches=True,
            demand=demand,
        )

        assert result == output_path
        written = pd.read_excel(output_path, sheet_name="LEAP", header=2)
        assert not written["Branch Path"].str.endswith(r"\Wind").any()
        assert set(written["Variable"]) == {
            aggregated_demand_workflow.ACTIVITY_VARIABLE_NAME,
            aggregated_demand_workflow.INTENSITY_VARIABLE_NAME,
        }
        pd.testing.assert_frame_equal(demand, original)
        assert demand["value"].sum() == pytest.approx(9.0, abs=1e-12)

    def test_all_zero_branch_is_retained_only_when_present_in_template(self, tmp_path):
        demand = pd.DataFrame(
            [
                self._row("Road", "Wind", "Current Accounts", 2022, 0.0),
                self._row("Road", "Wind", "Reference", 2023, 0.0),
                self._row("Road", "Wind", "Target", 2023, 0.0),
                self._row("Road", "Solar", "Current Accounts", 2022, 0.0),
                self._row("Road", "Solar", "Reference", 2023, 0.0),
                self._row("Road", "Solar", "Target", 2023, 0.0),
            ]
        )
        template_path = tmp_path / "template.xlsx"
        pd.DataFrame(
            [
                {
                    "BranchID": 101,
                    "VariableID": variable_index,
                    "ScenarioID": scenario_index,
                    "Branch Path": r"Demand\All demand aggregated\Road\Wind",
                    "Variable": variable,
                    "Scenario": scenario,
                }
                for scenario_index, scenario in enumerate(
                    ["Current Accounts", "Reference", "Target"], start=1
                )
                for variable_index, variable in enumerate(
                    [
                        aggregated_demand_workflow.ACTIVITY_VARIABLE_NAME,
                        aggregated_demand_workflow.INTENSITY_VARIABLE_NAME,
                    ],
                    start=1,
                )
            ]
        ).to_excel(template_path, index=False, startrow=2)

        output_path = tmp_path / "aggregated_demand.xlsx"
        aggregated_demand_workflow.save_aggregated_demand_as_leap_workbook(
            economy="20_USA",
            output_path=output_path,
            scenarios=["Current Accounts", "Reference", "Target"],
            id_lookup_path=template_path,
            use_sector_branches=True,
            demand=demand,
        )

        written = pd.read_excel(output_path, sheet_name="LEAP", header=2)
        wind_rows = written[
            written["Branch Path"].eq(r"Demand\All demand aggregated\Road\Wind")
        ]
        assert not wind_rows.empty
        assert set(wind_rows["Scenario"]) == {"Current Accounts", "Reference", "Target"}
        viewing = pd.read_excel(output_path, sheet_name="FOR_VIEWING", header=2)
        wind_activity = viewing[
            viewing["Branch Path"].eq(r"Demand\All demand aggregated\Road\Wind")
            & viewing["Variable"].eq(aggregated_demand_workflow.ACTIVITY_VARIABLE_NAME)
        ]
        current_accounts = wind_activity[
            wind_activity["Scenario"].eq("Current Accounts")
        ]
        projection_rows = wind_activity[
            wind_activity["Scenario"].isin(["Reference", "Target"])
        ]
        assert current_accounts["2022"].eq(0.0).all()
        projection_year_columns = [str(year) for year in range(2023, 2061)]
        assert projection_rows[projection_year_columns].eq(0.0).all().all()
        assert not written["Branch Path"].str.endswith(r"\Solar").any()

    def test_any_nonzero_value_retains_all_requested_scenarios(self, tmp_path):
        demand = pd.DataFrame(
            [
                self._row("Road", "Hydrogen", "Reference", 2023, 0.0),
                self._row("Road", "Hydrogen", "Reference", 2024, 1.0),
                self._row("Road", "Hydrogen", "Target", 2023, 0.0),
                self._row("Road", "Hydrogen", "Target", 2024, 0.0),
            ]
        )
        output_path = tmp_path / "aggregated_demand.xlsx"

        aggregated_demand_workflow.save_aggregated_demand_as_leap_workbook(
            economy="20_USA",
            output_path=output_path,
            scenarios=["Current Accounts", "Reference", "Target"],
            id_lookup_path=None,
            use_sector_branches=True,
            demand=demand,
        )

        written = pd.read_excel(output_path, sheet_name="LEAP", header=2)
        branch_rows = written[
            written["Branch Path"].eq(
                r"Demand\All demand aggregated\Road\Hydrogen"
            )
        ]
        assert set(branch_rows["Scenario"]) == {
            "Current Accounts",
            "Reference",
            "Target",
        }
        assert len(branch_rows) == 6
        current_accounts = branch_rows[
            branch_rows["Scenario"].eq("Current Accounts")
        ]
        assert set(current_accounts["Expression"].astype(str)) == {"0", "1"}


class TestContributionsSheet:
    """D5B.3/D5B.4: an opt-in Contributions sheet recording per-branch provenance
    behind each aggregated (fuel, scenario, year) total, and a reconciliation
    check that it sums back exactly to the aggregate it explains."""

    def _demand_and_contributions(self):
        demand = pd.DataFrame(
            [
                {"economy": "20_USA", "scenario": "Reference", "leap_fuel_name": "Electricity", "year": 2022, "value": 10.0},
                {"economy": "20_USA", "scenario": "Reference", "leap_fuel_name": "Electricity", "year": 2023, "value": 12.0},
            ]
        )
        contributions = pd.DataFrame(
            [
                {
                    "economy": "20_USA", "scenario": "Reference", "year": 2022,
                    "source_system": "NINTH", "source_sector_or_flow": "14_industry_sector",
                    "source_fuel_or_product": "electricity", "leap_fuel_name": "Electricity",
                    "allocation_method": "direct", "allocation_share": 0.6, "allocated_value": 6.0,
                },
                {
                    "economy": "20_USA", "scenario": "Reference", "year": 2022,
                    "source_system": "NINTH", "source_sector_or_flow": "16_01_buildings",
                    "source_fuel_or_product": "electricity", "leap_fuel_name": "Electricity",
                    "allocation_method": "direct", "allocation_share": 0.4, "allocated_value": 4.0,
                },
                {
                    "economy": "20_USA", "scenario": "Reference", "year": 2023,
                    "source_system": "NINTH", "source_sector_or_flow": "14_industry_sector",
                    "source_fuel_or_product": "electricity", "leap_fuel_name": "Electricity",
                    "allocation_method": "direct", "allocation_share": 1.0, "allocated_value": 12.0,
                },
            ]
        )
        return demand, contributions

    def test_reconciled_contributions_produce_no_warning(self, capsys):
        demand, contributions = self._demand_and_contributions()
        aggregated_demand_workflow._warn_contributions_do_not_reconcile(demand, contributions)
        assert "WARN" not in capsys.readouterr().out

    def test_mismatched_contributions_warn(self, capsys):
        demand, contributions = self._demand_and_contributions()
        contributions.loc[0, "allocated_value"] = 6.5  # 2022 total now 10.5, not 10.0
        aggregated_demand_workflow._warn_contributions_do_not_reconcile(demand, contributions)
        assert "WARN" in capsys.readouterr().out

    def test_raw_contribution_sign_is_preserved_separately(self):
        provenance = pd.DataFrame({"allocated_value": [-4.5, 2.0]})
        provenance = aggregated_demand_workflow._format_contribution_demand_magnitudes(
            provenance
        )

        assert provenance["allocated_value"].tolist() == [4.5, 2.0]
        assert provenance["raw_allocated_value"].tolist() == [-4.5, 2.0]

    def test_save_workbook_writes_contributions_sheet_when_opted_in(self, tmp_path, monkeypatch):
        demand, contributions = self._demand_and_contributions()

        def _fake_build_all_scenarios(*args, return_provenance=False, **kwargs):
            if return_provenance:
                return demand.copy(), contributions.copy()
            return demand.copy()

        monkeypatch.setattr(
            aggregated_demand_workflow,
            "build_aggregated_demand_all_scenarios",
            _fake_build_all_scenarios,
        )
        out_path = tmp_path / "aggregated_demand_20_USA.xlsx"
        result = aggregated_demand_workflow.save_aggregated_demand_as_leap_workbook(
            economy="20_USA",
            output_path=out_path,
            id_lookup_path=None,
            write_contributions=True,
        )
        assert result == out_path
        written = pd.read_excel(out_path, sheet_name="Contributions")
        assert set(written["source_sector_or_flow"]) == {"14_industry_sector", "16_01_buildings"}
        assert written["allocated_value"].sum() == pytest.approx(22.0)

    def test_save_workbook_omits_contributions_sheet_by_default(self, tmp_path, monkeypatch):
        demand, contributions = self._demand_and_contributions()

        def _fake_build_all_scenarios(*args, return_provenance=False, **kwargs):
            if return_provenance:
                return demand.copy(), contributions.copy()
            return demand.copy()

        monkeypatch.setattr(
            aggregated_demand_workflow,
            "build_aggregated_demand_all_scenarios",
            _fake_build_all_scenarios,
        )
        out_path = tmp_path / "aggregated_demand_20_USA.xlsx"
        aggregated_demand_workflow.save_aggregated_demand_as_leap_workbook(
            economy="20_USA",
            output_path=out_path,
            id_lookup_path=None,
        )
        import openpyxl
        sheet_names = openpyxl.load_workbook(out_path, read_only=True).sheetnames
        assert "Contributions" not in sheet_names
